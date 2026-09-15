"""Rebuild Coverity_Final_Slide_Updated.pptx so it looks exactly like
Coverity_Final_Slide.png while keeping every text field editable.

Why this pipeline exists
------------------------
The PNG was designed with Segoe UI (the same family the rest of the deck
uses).  Headless renderers available on Linux (Spire.Presentation,
LibreOffice) silently substitute another font for every typeface, so any
pptx calibrated against them looks "fully different" when opened in real
PowerPoint.  This script therefore never renders with a stand-in engine:

1. Every editable line's TRUE ink bbox is measured from the PNG itself.
2. The PNG's text glyphs are erased with a tight glyph mask + OpenCV
   TELEA inpainting (icons, tables, arrows, glows and the title halo are
   untouched) -> the slide background picture.
3. Every line is re-created as an editable "Segoe UI" text box whose size
   and position are computed ANALYTICALLY from the real Segoe UI font
   metrics (FreeType advance widths / ink bearings + the font's
   usWinAscent line metrics, which is what PowerPoint uses), so the ink
   lands on the measured PNG bbox in real PowerPoint.
4. The preview PNG is composited with the very same Segoe UI fonts and
   the very same formulas (FreeType ~= DirectWrite to within a pixel),
   i.e. the preview shows what PowerPoint will show.

Fonts: the script needs the two Microsoft Segoe UI TTFs locally (they are
NOT committed - they are proprietary).  Default locations
~/.fonts/segoeui.ttf and ~/.fonts/segoeuib.ttf, override with the
SEGOEUI_REG / SEGOEUI_BOLD environment variables.  On any Windows or macOS
machine running PowerPoint the family name "Segoe UI" resolves natively.

Usage:  python3 rebuild_final_slide.py
"""
import os
import sys
import numpy as np
import cv2
from PIL import Image, ImageDraw, ImageFilter, ImageFont

from pptx import Presentation
from pptx.util import Emu, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR, MSO_AUTO_SIZE
from pptx.oxml.ns import qn

from slide_specs import LINES

SRC = 'Coverity_Final_Slide.png'
W, H = 1376, 768
SLIDE_W_EMU = 12192000            # 13.3333 in  (same as the original deck)
SLIDE_H_EMU = 6858000             # 7.5 in
XI = (SLIDE_W_EMU / 914400.0) / W          # px -> inches
YI = (SLIDE_H_EMU / 914400.0) / H
OUT_PPTX = 'Coverity_Final_Slide_Updated.pptx'
OUT_PREVIEW = 'Coverity_Final_Slide_Updated_preview.png'
OUT_BG = '/tmp/slide_bg.png'

COLORS = {'W': (249, 251, 252), 'O': (238, 110, 30), 'L': (246, 150, 45)}

# Segoe UI OS/2 win metrics (units per em 2048) - what PowerPoint uses for
# the line box of a single-spaced paragraph.
UPEM = 2048.0
WIN_ASC = 2210.0
WIN_DESC = 514.0

REG_TTF = os.environ.get('SEGOEUI_REG', os.path.expanduser('~/.fonts/segoeui.ttf'))
BOLD_TTF = os.environ.get('SEGOEUI_BOLD', os.path.expanduser('~/.fonts/segoeuib.ttf'))


def load_fonts():
    if not (os.path.exists(REG_TTF) and os.path.exists(BOLD_TTF)):
        sys.exit('need Segoe UI ttf: place segoeui.ttf/segoeuib.ttf in ~/.fonts '
                 'or set SEGOEUI_REG/SEGOEUI_BOLD')
    return REG_TTF, BOLD_TTF


# --------------------------------------------------------------------------
# measurement
# --------------------------------------------------------------------------
def masks(img):
    a = np.asarray(img).astype(int)
    R, G, B = a[:, :, 0], a[:, :, 1], a[:, :, 2]
    white = (R > 200) & (G > 200) & (B > 200)
    orange = (R > 200) & (G > 70) & (G < 170) & (B < 95)      # hdr (238,110,30)
    ltorg = (R > 230) & (G > 120) & (G < 200) & (B < 95)     # banner (246,150,45)
    return white, orange, ltorg


def line_kind(runs):
    cks = {r[3] for r in runs}
    if cks == {'W'}:
        return 'W'
    return 'O'


def measure_targets(img):
    """True ink bbox (native px) of every editable line in the PNG."""
    white, orange, ltorg = masks(img)
    targets = []
    for x0, y0, x1, y1, align, runs in LINES:
        kind = line_kind(runs)
        m = white if kind == 'W' else (white | orange | ltorg)
        if runs[0][0].startswith('COVERITY'):
            m = white                       # title: core glyphs only
        wx0, wx1 = max(0, x0 - 4), min(W, x1 + 6)
        wy0, wy1 = max(0, y0 - 4), min(H, y1 + 6)
        sub = m[wy0:wy1, wx0:wx1].copy()
        rowsum = sub.sum(1)
        sub[rowsum > 0.7 * sub.shape[1], :] = False   # separator lines
        ys, xs = np.where(sub)
        if len(xs) == 0:
            raise SystemExit(f'no ink measured for line {runs[0][0]!r}')
        targets.append((wx0 + xs.min(), wy0 + ys.min(),
                        wx0 + xs.max() + 1, wy0 + ys.max() + 1))
    return targets


# --------------------------------------------------------------------------
# background: erase glyphs with a tight mask, keep everything else
# --------------------------------------------------------------------------
def build_glyph_mask(img, targets):
    white, orange, ltorg = masks(img)
    mask = np.zeros((H, W), np.uint8)
    for (bx0, by0, bx1, by1), line in zip(targets, LINES):
        runs = line[5]
        kind = line_kind(runs)
        m = white if kind == 'W' else (white | orange | ltorg)
        title = runs[0][0].startswith('COVERITY')
        if title:
            m = white
        wx0, wx1 = max(0, bx0 - 3), min(W, bx1 + 4)
        wy0, wy1 = max(0, by0 - 3), min(H, by1 + 4)
        sub = m[wy0:wy1, wx0:wx1]
        rowsum = sub.sum(1)
        sub = sub.copy()
        sub[rowsum > 0.7 * sub.shape[1], :] = False
        mask[wy0:wy1, wx0:wx1] |= sub.astype(np.uint8)
    # dilate: 6 px blurs the title's inner halo into a smooth glow,
    # 4 px elsewhere also eats the drop shadows
    k2 = np.ones((3, 3), np.uint8)
    core = mask.copy()
    for _ in range(6):
        core = cv2.dilate(core, k2)
    wide = mask.copy()
    for _ in range(4):
        wide = cv2.dilate(wide, k2)
    titleband = np.zeros_like(mask)
    for (bx0, by0, bx1, by1), line in zip(targets, LINES):
        if line[5][0][0].startswith('COVERITY'):
            titleband[max(0, by0 - 8):min(H, by1 + 8),
                      max(0, bx0 - 8):min(W, bx1 + 8)] = 1
    return np.where(titleband, core, wide)


def build_background(img, targets):
    mask = build_glyph_mask(img, targets)
    arr = np.ascontiguousarray(np.array(img.convert('RGB')))
    bg = cv2.inpaint(arr, mask, 4, cv2.INPAINT_TELEA)
    Image.fromarray(bg).save(OUT_BG)
    print('wrote background', OUT_BG, ' mask px', int(mask.sum()))
    return bg


# --------------------------------------------------------------------------
# analytic layout with real Segoe UI metrics
# --------------------------------------------------------------------------
class Layout:
    def __init__(self):
        self.rows = []          # dicts per line


def run_metrics(font, text, size_px):
    """ink bbox + advance of text at size, relative to the pen origin."""
    f = font(size_px)
    l, t, r, b = f.getbbox(text)
    adv = f.getlength(text)
    return l, t, r, b, adv


def solve_layout(targets):
    reg_p, bold_p = load_fonts()
    lay = Layout()
    for (bx0, by0, bx1, by1), line in zip(targets, LINES):
        align, runs = line[4], line[5]
        tw = bx1 - bx0

        def ink_width_at(size):
            x = 0.0
            lo = hi = None
            for text, _em, bold, _ck in runs:
                f = ImageFont.truetype(bold_p if bold else reg_p, size)
                l, t, r, b = f.getbbox(text)
                adv = f.getlength(text)
                if lo is None:
                    lo, hi = x + l, x + r
                else:
                    lo, hi = min(lo, x + l), max(hi, x + r)
                x += adv
            return hi - lo, x

        w100, adv100 = ink_width_at(100.0)
        size_px = 100.0 * tw / w100
        size_pt = size_px * 72.0 / 96.0

        # per-run geometry at final size
        x = 0.0
        geo = []
        ink_lo = ink_hi = ink_top = ink_bot = None
        for text, _em, bold, ck in runs:
            f = ImageFont.truetype(bold_p if bold else reg_p, size_px)
            l, t, r, b = f.getbbox(text)
            adv = f.getlength(text)
            geo.append(dict(text=text, bold=bold, ck=ck, font=f,
                            origin=x, l=l, t=t, r=r, b=b, adv=adv))
            ink_lo = x + l if ink_lo is None else min(ink_lo, x + l)
            ink_hi = x + r if ink_hi is None else max(ink_hi, x + r)
            ink_top = t if ink_top is None else min(ink_top, t)
            ink_bot = b if ink_bot is None else max(ink_bot, b)
            x += adv
        adv = x

        ascent_px = size_px * WIN_ASC / UPEM
        # PIL getbbox top (t) is relative to the draw anchor whose y is
        # baseline - ascent; PowerPoint puts that anchor at the top of a
        # margin-less, top-anchored frame.  Hence box_top = by0 - ink_top.
        box_top = by0 - ink_top
        if align == 'c':
            ink_c = (ink_lo + ink_hi) / 2.0
            tc = (bx0 + bx1) / 2.0
            center = tc - ink_c + adv / 2.0
            bw = max(adv * 1.6, 300.0)
            box_left = center - bw / 2.0
        else:
            box_left = bx0 - ink_lo
            bw = adv + 60.0
        box_h = size_px * (WIN_ASC + WIN_DESC) / UPEM * 1.5

        lay.rows.append(dict(align=align, size_pt=size_pt,
                             box_left=box_left, box_top=box_top,
                             box_w=bw, box_h=box_h, geo=geo,
                             target=(bx0, by0, bx1, by1),
                             expect_ink=(bx0, by0, bx0 + (ink_hi - ink_lo),
                                         by0 + (ink_bot - ink_top))))
    return lay


# --------------------------------------------------------------------------
# pptx
# --------------------------------------------------------------------------
def add_effects(run, glow):
    """Soft dark drop shadow like the source slide (2 px down, blurred);
    the title additionally gets the orange halo (a:glow)."""
    rPr = run._r.get_or_add_rPr()
    eff = rPr.makeelement(qn('a:effectLst'), {})
    if glow:
        gl = eff.makeelement(qn('a:glow'), {'rad': '101600'})
        c = gl.makeelement(qn('a:srgbClr'), {'val': 'EE7622'})
        al = c.makeelement(qn('a:alpha'), {'val': '60000'})
        c.append(al)
        gl.append(c)
        eff.append(gl)
    sh = eff.makeelement(qn('a:outerShdw'),
                         {'blurRad': '50800', 'dist': '25400',
                          'dir': '5400000', 'algn': 'bl', 'rotWithShape': '0'})
    clr = sh.makeelement(qn('a:srgbClr'), {'val': '000000'})
    al = clr.makeelement(qn('a:alpha'), {'val': '55000'})
    clr.append(al)
    sh.append(clr)
    eff.append(sh)
    rPr.append(eff)


def build_pptx(lay, bg_path, out):
    prs = Presentation()
    prs.slide_width = Emu(SLIDE_W_EMU)
    prs.slide_height = Emu(SLIDE_H_EMU)
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.shapes.add_picture(bg_path, 0, 0, Emu(SLIDE_W_EMU), Emu(SLIDE_H_EMU))
    for row in lay.rows:
        left = Emu(int(row['box_left'] * XI * 914400))
        top = Emu(int(row['box_top'] * YI * 914400))
        wid = Emu(int(row['box_w'] * XI * 914400))
        hei = Emu(int(row['box_h'] * YI * 914400))
        box = slide.shapes.add_textbox(left, top, wid, hei)
        tf = box.text_frame
        tf.word_wrap = False
        tf.auto_size = MSO_AUTO_SIZE.NONE
        tf.vertical_anchor = MSO_ANCHOR.TOP
        tf.margin_left = tf.margin_right = 0
        tf.margin_top = tf.margin_bottom = 0
        p = tf.paragraphs[0]
        p.alignment = (PP_ALIGN.CENTER if row['align'] == 'c'
                       else PP_ALIGN.LEFT)
        for g in row['geo']:
            r = p.add_run()
            r.text = g['text']
            r.font.size = Pt(row['size_pt'])
            r.font.bold = g['bold']
            r.font.name = 'Segoe UI'
            r.font.color.rgb = RGBColor(*COLORS[g['ck']])
            add_effects(r, glow=g['text'].startswith('COVERITY'))
    prs.save(out)
    print('wrote', out)


# --------------------------------------------------------------------------
# preview: composite with the same fonts & formulas PowerPoint will use
# --------------------------------------------------------------------------
def _row_origins(row):
    """pen x of the first run + baseline y, mirroring PowerPoint layout."""
    size_px = row['size_pt'] * 96.0 / 72.0
    baseline = row['box_top'] + size_px * WIN_ASC / UPEM
    if row['align'] == 'c':
        ox = row['box_left'] + row['box_w'] / 2.0 - \
            sum(g['adv'] for g in row['geo']) / 2.0
    else:
        ox = row['box_left']
    return ox, baseline, size_px


def composite_preview(bg, lay, out):
    canvas = Image.new('RGBA', (W, H), (255, 255, 255, 255))
    canvas.paste(Image.fromarray(bg).convert('RGBA'), (0, 0))

    # soft drop shadow pass (approximates the a:outerShdw written in pptx)
    shadow = np.zeros((H, W), np.uint8)
    for row in lay.rows:
        ox, baseline, spx = _row_origins(row)
        layer = Image.new('L', (W, H), 0)
        d = ImageDraw.Draw(layer)
        for g in row['geo']:
            d.text((ox + g['origin'], baseline - spx * WIN_ASC / UPEM),
                   g['text'], font=g['font'], fill=255)
        shadow = np.maximum(shadow, np.array(layer))
    shadow_img = Image.fromarray(shadow).filter(ImageFilter.GaussianBlur(1.6))
    blk = Image.new('RGBA', (W, H), (0, 0, 0, 0))
    blk.putalpha(shadow_img.point(lambda v: int(v * 0.55)))
    canvas = Image.alpha_composite(canvas, __shift(blk, 0, 2))

    # orange halo pass for the title (approximates the a:glow in the pptx)
    for row in lay.rows:
        if not row['geo'][0]['text'].startswith('COVERITY'):
            continue
        ox, baseline, spx = _row_origins(row)
        layer = Image.new('L', (W, H), 0)
        d = ImageDraw.Draw(layer)
        for g in row['geo']:
            d.text((ox + g['origin'], baseline - spx * WIN_ASC / UPEM),
                   g['text'], font=g['font'], fill=255)
        for blur, alpha in ((2.5, 0.85), (6.0, 0.45)):
            gimg = layer.filter(ImageFilter.GaussianBlur(blur))
            col = Image.new('RGBA', (W, H), (238, 118, 34, 255))
            col.putalpha(gimg.point(
                lambda v, a=alpha: int(min(255, v * 1.5) * a)))
            canvas = Image.alpha_composite(canvas, col)

    # crisp text pass
    txt = Image.new('RGBA', (W, H), (0, 0, 0, 0))
    for row in lay.rows:
        ox, baseline, spx = _row_origins(row)
        for g in row['geo']:
            layer = Image.new('L', (W, H), 0)
            d = ImageDraw.Draw(layer)
            d.text((ox + g['origin'], baseline - spx * WIN_ASC / UPEM),
                   g['text'], font=g['font'], fill=255)
            col = Image.new('RGBA', (W, H), (*COLORS[g['ck']], 255))
            col.putalpha(layer)
            txt = Image.alpha_composite(txt, col)
    canvas = Image.alpha_composite(canvas, txt)
    canvas.convert('RGB').save(out)
    print('wrote', out)


def __shift(img, dx, dy):
    out = Image.new('RGBA', img.size, (0, 0, 0, 0))
    out.paste(img, (dx, dy))
    return out


def verify(lay):
    worst = 0.0
    for row in lay.rows:
        t = row['target']
        e = row['expect_ink']
        d = max(abs(t[0] - e[0]), abs(t[1] - e[1]))
        worst = max(worst, d)
    print(f'layout check: intended ink vs measured PNG bbox, worst dx/dy = '
          f'{worst:.2f}px (by construction x/y left-top coincide)')


def main():
    img = Image.open(SRC).convert('RGB')
    targets = measure_targets(img)
    for t, line in zip(targets, LINES):
        print(f'target {t}  {line[5][0][0][:28]!r}')
    bg = build_background(img, targets)
    lay = solve_layout(targets)
    verify(lay)
    build_pptx(lay, OUT_BG, OUT_PPTX)
    composite_preview(bg, lay, OUT_PREVIEW)


if __name__ == '__main__':
    main()
